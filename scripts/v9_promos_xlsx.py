#!/usr/bin/env python3
"""
v9 — xlsx-таблица 3000 промокодов закрытого запуска no-reality.
Источник: scripts/v9_promos_dump.json (выгрузка из Supabase после
вставки). Дизайн — по дизайн-системе xlsx-скилла (templates/base.py):
B2-canvas, primary-заголовок, зебра-строки, без сетки, итоговая строка.
"""
import json
import os
import sys
from datetime import datetime

XLSX_SKILL_DIR = "/home/z/my-project/skills/xlsx"
for sub in [XLSX_SKILL_DIR, os.path.join(XLSX_SKILL_DIR, "templates")]:
    if sub not in sys.path:
        sys.path.insert(0, sub)

from base import (  # noqa: E402
    setup_sheet,
    style_header_row,
    style_data_row,
    style_total_row,
    font_caption,
    auto_fit_columns,
)
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment  # noqa: E402

DUMP = "/home/z/my-project/scripts/v9_promos_dump.json"
OUT_XLSX = "/home/z/my-project/download/no-reality-promo-codes-3000.xlsx"
OUT_CSV = "/home/z/my-project/download/no-reality-promo-codes-3000.csv"

with open(DUMP, "r", encoding="utf-8") as f:
    dump = json.load(f)

codes = dump["codes"]
assert len(codes) == 3000, f"ожидалось 3000 кодов, получено {len(codes)}"
flat_set = {c["flat"] for c in codes}
assert len(flat_set) == 3000, "коды не уникальны!"

headers = ["№", "Промокод", "Партия", "Статус", "Создан (UTC)"]
last_col = len(headers) + 1  # B..F → 6

wb = Workbook()
ws = wb.active
ws.title = "Промокоды"

setup_sheet(
    ws,
    title="no-reality. — промокоды закрытого запуска (3000 шт.)",
    last_col=last_col,
)

# подзаголовок в строке 3 (между заголовком и шапкой)
ws.cell(
    row=3,
    column=2,
    value=f"партия {dump['batch']} · выгружено {dump['generatedAt'][:10]} · "
          "ввод на no-reality.fun/auth — код одноразовый",
)
ws.cell(row=3, column=2).font = font_caption()
ws.row_dimensions[3].height = 14

# шапка (row 4)
for col_idx, h in enumerate(headers, start=2):
    ws.cell(row=4, column=col_idx, value=h)
style_header_row(ws, row_num=4, col_start=2, col_end=last_col)

# данные (row 5+)
for i, c in enumerate(codes):
    row_num = 5 + i
    ws.cell(row=row_num, column=2, value=c["n"])
    ws.cell(row=row_num, column=3, value=c["pretty"])
    ws.cell(row=row_num, column=4, value=dump["batch"])
    ws.cell(row=row_num, column=5, value="не использован")
    created = str(c["createdAt"])[:16].replace("T", " ")
    ws.cell(row=row_num, column=6, value=created)
    style_data_row(ws, row_num=row_num, col_start=2, col_end=last_col, row_index=i)
    # выравнивание: № вправо, код моно-стилем по центру не надо — текст влево
    ws.cell(row=row_num, column=2).alignment = Alignment(horizontal="right", vertical="center")

last_data_row = 4 + len(codes)

# итоговая строка
total_row = last_data_row + 1
ws.cell(row=total_row, column=2, value="Итого")
ws.cell(row=total_row, column=3, value=f"{len(codes)} шт. · уникальных: {len(flat_set)}")
ws.cell(row=total_row, column=4, value=dump["batch"])
ws.cell(row=total_row, column=5, value="не использован")
ws.cell(row=total_row, column=6, value="")
style_total_row(ws, row_num=total_row, col_start=2, col_end=last_col)

# примечание
note_row = total_row + 2
ws.cell(
    row=note_row,
    column=2,
    value="Как работает: регистрация на /auth с кодом = мгновенный аккаунт; без кода — лист ожидания. "
          "Код гасится атомарно при первой активации, повторно не применяется. "
          "Перебор невозможен: 31^12 комбинаций + бюджет промахов (8/час, 20/сутки на IP) на API.",
)
ws.cell(row=note_row, column=2).font = font_caption()

ws.freeze_panes = "B5"
auto_fit_columns(ws, min_width=8, max_width=34, header_row=4, data_start_row=5)

wb.properties.creator = "Z.ai"
os.makedirs(os.path.dirname(OUT_XLSX), exist_ok=True)
wb.save(OUT_XLSX)
print(f"saved {OUT_XLSX}")

# CSV-копия для импорта куда угодно
import csv  # noqa: E402

with open(OUT_CSV, "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(["n", "code", "batch", "status", "created_utc"])
    for c in codes:
        w.writerow([c["n"], c["pretty"], dump["batch"], "не использован", str(c["createdAt"])[:19].replace("T", " ")])
print(f"saved {OUT_CSV}")
